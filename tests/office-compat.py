#!/usr/bin/env python3
"""Optional interoperability validation with python-docx and openpyxl.
Run: uv run --with python-docx --with openpyxl python tests/office-compat.py
The production helper itself has no third-party dependencies.
"""
import importlib.util
from pathlib import Path
import tempfile

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('office_helper',ROOT/'crates/buddy-agent/src/tools/office.py')
helper=importlib.util.module_from_spec(spec)
spec.loader.exec_module(helper)

with tempfile.TemporaryDirectory(prefix='drops-office-compat-') as directory:
    base=Path(directory)
    doc=Document()
    paragraph=doc.add_paragraph()
    paragraph.add_run('Hello ').bold=True
    paragraph.add_run('coworker!').italic=True
    doc.add_table(rows=1,cols=1).cell(0,0).text='Table preserved'
    source=base/'report.docx';output=base/'report-edited.docx'
    doc.save(source)
    read=helper.run(dict(_root=str(base),path=str(source)))
    helper.run(dict(_root=str(base),_edit=True,path=str(source),output_path=str(output),expected_sha256=read['sha256'],replacements=[dict(old_text='lo cowork',new_text='pful work')]))
    updated=Document(output)
    assert updated.paragraphs[0].text=='Helpful worker!'
    assert updated.paragraphs[0].runs[0].bold
    assert updated.paragraphs[0].runs[1].italic
    assert updated.tables[0].cell(0,0).text=='Table preserved'
    assert Document(source).paragraphs[0].text=='Hello coworker!'
    book=Workbook();sheet=book.active;sheet.title='Budget'
    sheet['A1']='Revenue';sheet['A1'].font=Font(bold=True)
    sheet['B1']=10;sheet['B1'].number_format='$#,##0.00'
    sheet['C1']='=B1*2'
    book.create_sheet('Other')['A1']='Unchanged'
    source=base/'budget.xlsx';output=base/'budget-edited.xlsx';book.save(source)
    read=helper.run(dict(_root=str(base),path=str(source)))
    helper.run(dict(_root=str(base),_edit=True,path=str(source),output_path=str(output),expected_sha256=read['sha256'],sheet='Budget',cells=[dict(cell='A1',value='Updated revenue'),dict(cell='B1',value=30)]))
    updated=load_workbook(output)
    assert updated['Budget']['A1'].value=='Updated revenue'
    assert updated['Budget']['A1'].font.bold
    assert updated['Budget']['B1'].value==30
    assert updated['Budget']['B1'].number_format=='$#,##0.00'
    assert updated['Budget']['C1'].value=='=B1*2'
    assert updated['Other']['A1'].value=='Unchanged'
    assert updated.calculation.fullCalcOnLoad
    assert load_workbook(source)['Budget']['B1'].value==10
print('PASS: real DOCX/XLSX packages reopen in independent readers; Word formatting/tables and Excel styles/formulas/other sheets survive edits')
